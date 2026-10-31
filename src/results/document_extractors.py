"""Format-aware extraction helpers for exchange financial-result attachments.

Metrics are only promoted when the source contains tagged/structured facts. PDF text,
OCR-like text, and arbitrary HTML tables remain review material rather than silently
being treated as verified financial metrics.
"""
import csv
from io import BytesIO, StringIO
import re
import xml.etree.ElementTree as ET
from typing import Any, Dict, Optional

from bs4 import BeautifulSoup
from openpyxl import load_workbook

MAX_DOCUMENT_BYTES = 15_000_000


def detect_document_format(content: bytes, content_type: str = "", url: str = "") -> str:
    """Identify common filing formats using signatures first and metadata second."""
    if not content:
        return "unknown"
    head = content[:512].lstrip().lower()
    mime = (content_type or "").lower()
    path = (url or "").split("?", 1)[0].lower()
    if content.startswith(b"%PDF-"):
        return "pdf"
    if content.startswith(b"PK\x03\x04"):
        return "xlsx" if path.endswith((".xlsx", ".xlsm")) or "spreadsheet" in mime or "excel" in mime else "office_zip"
    if head.startswith((b"<?xml", b"<xbrl", b"<html", b"<!doctype html")) or b"<html" in head:
        if b"nonfraction" in content[:200000].lower() or b"nonnumeric" in content[:200000].lower() or "ixbrl" in path:
            return "ixbrl"
        if b"<xbrl" in head or b":xbrl" in head:
            return "xbrl_xml"
        return "html"
    if path.endswith(".csv") or "text/csv" in mime:
        return "csv"
    if path.endswith((".xlsx", ".xls")) or "excel" in mime or "spreadsheet" in mime:
        return "xlsx"
    try:
        decoded = content[:4096].decode("utf-8-sig")
        if "," in decoded and "\n" in decoded and not "\x00" in decoded:
            return "csv"
    except UnicodeDecodeError:
        pass
    return "unknown"


def extract_ixbrl_as_xbrl(content: bytes) -> bytes:
    """Convert inline-XBRL facts/resources into a small XBRL instance for existing parsing.

    Keeps only contexts, units, and ix:nonFraction numeric facts. It does not evaluate
    JavaScript or external resources and rejects DTD/entity declarations.
    """
    if not content or len(content) > 10_000_000:
        raise ValueError("iXBRL content is empty or exceeds the 10 MB safety limit.")
    upper = content.upper()
    if b"<!DOCTYPE" in upper or b"<!ENTITY" in upper:
        raise ValueError("iXBRL document contains a prohibited DTD/entity declaration.")
    soup = BeautifulSoup(content, "lxml")
    resources = soup.find(lambda tag: getattr(tag, "name", "").lower().endswith("resources"))
    if resources is None:
        raise ValueError("No inline-XBRL resources block was found.")
    root = ET.Element("xbrl")
    for node in resources.find_all(recursive=True):
        name = node.name.split(":")[-1].lower()
        if name in {"context", "unit"}:
            element = ET.fromstring(str(node).encode("utf-8"))
            root.append(element)
    count = 0
    for fact in soup.find_all(lambda tag: getattr(tag, "name", "").lower().endswith("nonfraction")):
        fact_name = fact.get("name")
        context_ref = fact.get("contextref")
        if not fact_name or not context_ref:
            continue
        local_fact = fact_name.split(":")[-1]
        element = ET.SubElement(root, local_fact, {"contextRef": context_ref})
        unit_ref = fact.get("unitref")
        if unit_ref:
            element.set("unitRef", unit_ref)
        scale = fact.get("scale")
        sign = fact.get("sign")
        raw = re.sub(r"\s+", "", fact.get_text("", strip=True)).replace(",", "")
        if raw.startswith("(") and raw.endswith(")"):
            raw = "-" + raw[1:-1]
        if sign == "-":
            raw = raw[1:] if raw.startswith("-") else "-" + raw
        try:
            number = float(raw) * (10 ** int(scale or "0"))
        except (ValueError, OverflowError):
            continue
        element.text = str(number)
        count += 1
    if not count:
        raise ValueError("No usable inline-XBRL numeric facts were found.")
    return ET.tostring(root, encoding="utf-8", xml_declaration=True)


def extract_csv_rows(content: bytes) -> Dict[str, Any]:
    if not content or len(content) > MAX_DOCUMENT_BYTES:
        return {"status": "UNAVAILABLE", "reason": "CSV is empty or exceeds the size limit."}
    try:
        text = content.decode("utf-8-sig")
        rows = list(csv.reader(StringIO(text)))
    except (UnicodeDecodeError, csv.Error) as exc:
        return {"status": "UNAVAILABLE", "reason": f"CSV parsing failed ({type(exc).__name__})."}
    return {"status": "EXTRACTED" if rows else "UNAVAILABLE", "format": "csv", "rows": rows[:5000]}


def extract_xlsx_rows(content: bytes) -> Dict[str, Any]:
    if not content or len(content) > MAX_DOCUMENT_BYTES:
        return {"status": "UNAVAILABLE", "reason": "Workbook is empty or exceeds the size limit."}
    try:
        workbook = load_workbook(BytesIO(content), read_only=True, data_only=True)
        sheets = {}
        for sheet in workbook.worksheets[:20]:
            sheets[sheet.title] = [
                list(row) for row in sheet.iter_rows(max_row=5000, max_col=100, values_only=True)
            ]
        workbook.close()
        return {"status": "EXTRACTED" if sheets else "UNAVAILABLE", "format": "xlsx", "sheets": sheets}
    except Exception as exc:
        return {"status": "UNAVAILABLE", "reason": f"Workbook parsing failed ({type(exc).__name__})."}


def extract_html_tables(content: bytes) -> Dict[str, Any]:
    if not content or len(content) > MAX_DOCUMENT_BYTES:
        return {"status": "UNAVAILABLE", "reason": "HTML is empty or exceeds the size limit."}
    try:
        soup = BeautifulSoup(content, "lxml")
        tables = []
        for table in soup.find_all("table")[:100]:
            rows = []
            for tr in table.find_all("tr")[:5000]:
                rows.append([cell.get_text(" ", strip=True) for cell in tr.find_all(["th", "td"])[:100]])
            if rows:
                tables.append(rows)
        return {"status": "EXTRACTED" if tables else "UNAVAILABLE", "format": "html", "tables": tables}
    except Exception as exc:
        return {"status": "UNAVAILABLE", "reason": f"HTML parsing failed ({type(exc).__name__})."}


def extract_structured_document(content: bytes, content_type: str = "", url: str = "") -> Dict[str, Any]:
    """Return bounded structured content; caller must validate issuer, period, units and basis."""
    fmt = detect_document_format(content, content_type, url)
    if fmt == "ixbrl":
        return {"status": "EXTRACTED", "format": fmt, "xbrl_xml": extract_ixbrl_as_xbrl(content)}
    if fmt == "xbrl_xml":
        upper = content.upper()
        if b"<!DOCTYPE" in upper or b"<!ENTITY" in upper:
            return {"status": "UNAVAILABLE", "reason": "XML contains a prohibited DTD/entity declaration."}
        try:
            ET.fromstring(content)
            return {"status": "EXTRACTED", "format": fmt, "xbrl_xml": content}
        except ET.ParseError:
            return {"status": "UNAVAILABLE", "reason": "Malformed XML attachment."}
    if fmt == "xlsx":
        return extract_xlsx_rows(content)
    if fmt == "csv":
        return extract_csv_rows(content)
    if fmt == "html":
        return extract_html_tables(content)
    if fmt == "pdf":
        return {"status": "REVIEW_REQUIRED", "format": "pdf", "reason": "Use bounded PDF text extraction; scanned PDFs require OCR and human validation."}
    return {"status": "UNAVAILABLE", "reason": "Unsupported or unrecognized attachment format."}


def extract_screener_quarter_metrics(html: bytes, company_name: str, period_ended: str) -> Dict[str, Any]:
    """Best-effort fallback from Screener's quarterly table, requiring issuer and period match.

    Screener is secondary only: caller must invoke this after exchange extraction fails and
    must retain the source URL and FALLBACK status. Values are not treated as exchange-verified.
    """
    if not html or len(html) > MAX_DOCUMENT_BYTES:
        return {"status": "UNAVAILABLE", "reason": "Screener page is empty or exceeds the size limit."}
    soup = BeautifulSoup(html, "lxml")
    title = soup.title.get_text(" ", strip=True) if soup.title else ""
    expected = re.sub(r"[^a-z0-9]+", " ", company_name.lower()).strip().replace(" limited", " ltd")
    observed = re.sub(r"[^a-z0-9]+", " ", title.lower()).strip().replace(" limited", " ltd")
    if not expected or expected not in observed:
        return {"status": "UNAVAILABLE", "reason": "Screener issuer title did not match the configured company."}
    try:
        from datetime import datetime
        target = datetime.strptime(period_ended[:10], "%Y-%m-%d") if re.match(r"^\d{4}-\d{2}-\d{2}", period_ended) else datetime.strptime(period_ended, "%d-%b-%Y")
        month_alias = target.strftime("%b %Y").lower()
    except (TypeError, ValueError):
        return {"status": "UNAVAILABLE", "reason": "Reporting period could not be normalized for Screener matching."}
    for table in soup.find_all("table"):
        rows = [[cell.get_text(" ", strip=True) for cell in tr.find_all(["th", "td"])] for tr in table.find_all("tr")]
        if not rows:
            continue
        headers = [re.sub(r"\s+", " ", value).strip().lower() for value in rows[0]]
        col = next((i for i, header in enumerate(headers) if month_alias in header), None)
        if col is None:
            continue
        values = {}
        for row in rows[1:]:
            if len(row) <= col:
                continue
            label = re.sub(r"[^a-z0-9]+", " ", row[0].lower()).strip()
            raw = row[col].replace(",", "").replace("%", "").strip()
            try:
                number = float(raw)
            except ValueError:
                continue
            if label in {"sales", "revenue", "revenue from operations"}:
                values["revenue_cr"] = number
            elif label in {"net profit", "net profit loss", "profit after tax", "pat"}:
                values["pat_cr"] = number
            elif label in {"eps", "eps in rs", "eps rs"}:
                values["eps"] = number
        if values:
            return {
                "status": "EXTRACTED",
                "source_status": "FALLBACK",
                "source_name": "Screener.in",
                "source_url": "",
                "period_label": month_alias,
                "revenue_cr": values.get("revenue_cr"),
                "pat_cr": values.get("pat_cr"),
                "eps": values.get("eps"),
                "cfo_cr": None,
                "revenue_yoy_pct": None,
                "pat_yoy_pct": None,
                "revenue_qoq_pct": None,
                "pat_qoq_pct": None,
                "validation_note": "Secondary-source figures; independently verify against the official exchange filing.",
            }
    return {"status": "UNAVAILABLE", "reason": "No matching quarterly period and recognized metric rows were found on Screener."}
